"""
Utility functions for file parsing and document reading.

Provides functions to extract text from PDF, DOCX, and TXT files.
Designed for handling uploaded files in the Streamlit UI.
"""

import io
import logging
from typing import Optional

logger = logging.getLogger(__name__)

# Deferred imports for optional dependencies
pypdf_available = False
docx_available = False

try:
    import pypdf
    pypdf_available = True
except ImportError:
    logger.warning("pypdf is not installed. PDF parsing will not be available.")

try:
    import docx
    docx_available = True
except ImportError:
    logger.warning("python-docx is not installed. DOCX parsing will not be available.")


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract text from a PDF file in memory.

    Args:
        file_bytes: Binary contents of the PDF file.

    Returns:
        Extracted plain text.
    """
    if not pypdf_available:
        raise ImportError(
            "pypdf package is not installed. Cannot parse PDF files."
        )

    try:
        pdf_file = io.BytesIO(file_bytes)
        reader = pypdf.PdfReader(pdf_file)
        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n".join(text_parts)
    except Exception as e:
        logger.exception(f"Failed to parse PDF file: {e}")
        raise ValueError(f"Error reading PDF file: {e}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX file in memory.

    Args:
        file_bytes: Binary contents of the DOCX file.

    Returns:
        Extracted plain text.
    """
    if not docx_available:
        raise ImportError(
            "python-docx package is not installed. Cannot parse DOCX files."
        )

    try:
        docx_file = io.BytesIO(file_bytes)
        doc = docx.Document(docx_file)
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text_parts.append(paragraph.text)
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_parts.append(cell.text)
                        
        return "\n".join(text_parts)
    except Exception as e:
        logger.exception(f"Failed to parse DOCX file: {e}")
        raise ValueError(f"Error reading DOCX file: {e}")


def extract_text_from_file(uploaded_file) -> str:
    """
    Detect file type and extract text from a Streamlit UploadedFile object.

    Args:
        uploaded_file: UploadedFile object from streamlit.file_uploader.

    Returns:
        Extracted plain text string.
    """
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()
    # Reset read pointer for safety
    uploaded_file.seek(0)

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        # Fallback to plain text
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return file_bytes.decode("latin-1")
            except UnicodeDecodeError as e:
                raise ValueError(
                    f"Could not decode text file {uploaded_file.name}: {e}"
                )
